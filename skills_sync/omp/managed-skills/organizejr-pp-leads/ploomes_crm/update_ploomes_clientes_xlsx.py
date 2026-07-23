#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import urllib.parse
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


FIELD_ALIASES = {
    "Nome - Pessoa": ["Nome - Pessoa", "Nome Pessoa"],
    "Cargo - Pessoa": ["Cargo - Pessoa", "Cargo"],
    "Departamento - Pessoa": ["Departamento - Pessoa", "Departamento"],
    "CPF - Pessoa": ["CPF - Pessoa", "CPF"],
    "Data de nascimento - Pessoa": ["Data de nascimento - Pessoa", "Data de nascimento"],
    "Telefones - Pessoa": ["Telefones - Pessoa", "Telefones", "Telefone - Pessoa"],
    "E-mail - Pessoa": ["E-mail - Pessoa", "Email - Pessoa", "E-mail", "Email"],
    "Razão Social - Empresa": [
        "Razão Social - Empresa",
        "Razão social",
        "Razão social - Empresa",
        "Nome - Empresa (ou B2C)",
        "Nome - Empresa",
    ],
    "Segmento de atuação - Empresa": ["Segmento de atuação - Empresa", "Segmento"],
    "Origem - Empresa": ["Origem - Empresa", "Origem"],
    "CNPJ - Empresa": ["CNPJ - Empresa", "CNPJ"],
    "Marcadores - Empresa": ["Marcadores - Empresa", "Marcadores"],
    "Estado - Empresa": ["Estado - Empresa", "Estado"],
    "Cidade - Empresa": ["Cidade - Empresa", "Cidade"],
    "Site": ["Site", "Site - Empresa"],
    "E-mail do Responsável": ["E-mail do Responsável", "Responsável", "Email do Responsável"],
    "Relação": ["Relação", "Relacao"],
    "Observações": ["Observações", "Observacoes"],
    "Enviar ao Ploomes?": ["Enviar ao Ploomes?"],
    "Status Importação": ["Status Importação"],
    "Ploomes Id Empresa": ["Ploomes Id Empresa"],
    "Ploomes Id Pessoa": ["Ploomes Id Pessoa"],
    "Erro Importação": ["Erro Importação"],
    "Última sincronização": ["Última sincronização"],
}


CONTROL_HEADERS = {
    "Enviar ao Ploomes?",
    "Status Importação",
    "Ploomes Id Empresa",
    "Ploomes Id Pessoa",
    "Erro Importação",
    "Última sincronização",
}


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"\d{8,}")


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.lower() in {"nan", "none", "null"}:
        return ""
    return text


def digits(value: Any) -> str:
    return "".join(ch for ch in clean(value) if ch.isdigit())


def slug(text: str) -> str:
    base = clean(text).casefold()
    base = re.sub(r"[^a-z0-9]+", "-", base)
    return base.strip("-")


def normalize(text: Any) -> str:
    return re.sub(r"\s+", " ", clean(text).casefold())


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Atualiza Ploomes_Clientes.xlsx a partir de artefatos OrganizeJr.")
    parser.add_argument("--lead-table", help="CSV/XLSX com a tabela principal de leads.")
    parser.add_argument("--use-case-config", help="Diretório do recorte/ICP para descobrir lead_table_path e output_dir.")
    parser.add_argument("--workbook", required=True, help="Caminho da planilha Ploomes_Clientes.xlsx.")
    parser.add_argument("--output", help="Salvar em novo arquivo. Obrigatório se não usar --in-place.")
    parser.add_argument("--in-place", action="store_true", help="Sobrescreve a planilha de entrada.")
    parser.add_argument("--enrichment-dir", help="Diretório opcional com JSONs de enrich.")
    parser.add_argument("--limit", type=int, default=None, help="Processa apenas as primeiras N linhas do lead table.")
    parser.add_argument("--min-score", type=float, default=55.0, help="Score mínimo para marcar envio quando a classe não ajudar.")
    parser.add_argument("--json", action="store_true", help="Imprime resumo em JSON.")
    args = parser.parse_args()
    if not args.lead_table and not args.use_case_config:
        parser.error("use --lead-table ou --use-case-config")
    if args.in_place and args.output:
        parser.error("use --in-place ou --output, não ambos")
    if not args.in_place and not args.output:
        parser.error("use --output ou --in-place")
    return args


def resolve_paths(args: argparse.Namespace) -> tuple[Path, Path | None, Path | None]:
    lead_table = Path(args.lead_table).expanduser().resolve() if args.lead_table else None
    enrichment_dir = Path(args.enrichment_dir).expanduser().resolve() if args.enrichment_dir else None
    if args.use_case_config:
        cfg_path = Path(args.use_case_config).expanduser().resolve()
        if cfg_path.is_dir():
            if lead_table is None:
                csv_matches = sorted(cfg_path.glob("lead-table-*.csv"))
                xlsx_matches = sorted(cfg_path.glob("lead-table-*.xlsx"))
                candidates = csv_matches or xlsx_matches
                if candidates:
                    lead_table = candidates[0].resolve()
            if enrichment_dir is None:
                enrichment_dir = (cfg_path / "outputs" / "enrichment").resolve()
        else:
            cfg = json.loads(cfg_path.read_text())
            if lead_table is None:
                lead_table = (cfg_path.parent / cfg["lead_table_path"]).resolve()
            if enrichment_dir is None:
                enrichment_dir = (cfg_path.parent / cfg.get("output_dir", "")).resolve() if cfg.get("output_dir") else None
    if lead_table is None:
        raise SystemExit("lead table ausente")
    return lead_table, enrichment_dir, Path(args.workbook).expanduser().resolve()


def read_table(path: Path, limit: int | None) -> list[dict[str, str]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as fh:
            rows = list(csv.DictReader(fh))
    elif suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        headers = [clean(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        rows = []
        for values in ws.iter_rows(min_row=2, values_only=True):
            item = {headers[i]: clean(values[i]) for i in range(len(headers)) if headers[i]}
            if any(item.values()):
                rows.append(item)
    else:
        raise SystemExit(f"formato não suportado: {path}")
    if limit is not None:
        return rows[:limit]
    return rows


def find_header_row(ws) -> tuple[int, list[str]]:
    for row_idx in range(1, min(ws.max_row, 10) + 1):
        values = [clean(ws.cell(row=row_idx, column=col).value) for col in range(1, ws.max_column + 1)]
        if "Razão Social - Empresa" in values or "Nome - Empresa (ou B2C)" in values:
            return row_idx, values
    raise SystemExit("não encontrei cabeçalhos da planilha Ploomes")


def build_header_map(headers: list[str]) -> dict[str, int]:
    lookup = {header: idx + 1 for idx, header in enumerate(headers) if header}
    resolved = {}
    for canonical, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if alias in lookup:
                resolved[canonical] = lookup[alias]
                break
    return resolved


def parse_city_state(text: str) -> tuple[str, str]:
    value = clean(text)
    if "/" in value:
        city, state = value.rsplit("/", 1)
        return clean(city), clean(state).upper()
    return value, ""

def normalize_website(value: Any) -> str:
    raw = clean(value)
    if not raw:
        return ""
    if raw.casefold().startswith("sem "):
        return ""
    if len(raw) > 200:
        return ""
    parsed = urllib.parse.urlparse(raw)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        return ""
    if "google." in parsed.netloc.casefold():
        return ""
    return raw


def contact_parts(value: str) -> tuple[str, str, str]:
    text = clean(value)
    if not text or text.casefold().startswith("sem "):
        return "", "", ""
    if EMAIL_RE.match(text):
        return "", text, ""
    if PHONE_RE.search(digits(text)):
        return "", "", digits(text)
    return text, "", ""


def classify_send_flag(row: dict[str, str], min_score: float) -> str:
    label = normalize(row.get("classe"))
    if "benchmark" in label:
        return "NÃO"
    if label.startswith("a") or label.startswith("b"):
        return "SIM"
    try:
        return "SIM" if float(clean(row.get("score"))) >= min_score else "NÃO"
    except ValueError:
        return "NÃO"


def derive_segment(row: dict[str, str]) -> str:
    icp = normalize(row.get("ICP"))
    if "comunica" in icp and "ej" in icp:
        return "Empresa Júnior de Comunicação"
    if "ej" in icp or "empresa júnior" in icp:
        return "Empresa Júnior"
    camada = clean(row.get("camada"))
    return camada.split(",", 1)[0].strip() if camada else ""


def join_unique(parts: list[str], sep: str = "; ") -> str:
    values = []
    seen = set()
    for part in parts:
        value = clean(part)
        key = value.casefold()
        if not value or key in seen:
            continue
        seen.add(key)
        values.append(value)
    return sep.join(values)


def load_enrichment(row: dict[str, str], enrichment_dir: Path | None) -> tuple[dict[str, Any], str]:
    candidates: list[Path] = []
    explicit = clean(row.get("pp_leads_enrich"))
    if explicit and explicit.endswith(".json"):
        exp_path = Path(explicit)
        if not exp_path.is_absolute() and enrichment_dir:
            exp_path = enrichment_dir / exp_path.name
        candidates.append(exp_path)
    cnpj = digits(row.get("CNPJ"))
    name_slug = slug(row.get("lead", ""))
    if enrichment_dir and enrichment_dir.exists():
        if cnpj:
            candidates.extend([
                enrichment_dir / f"{cnpj}.json",
                enrichment_dir / f"{cnpj}-enrich.json",
                enrichment_dir / f"{cnpj}_enrich.json",
            ])
        if name_slug:
            candidates.extend([
                enrichment_dir / f"{name_slug}.json",
                enrichment_dir / f"{name_slug}-enrich.json",
                enrichment_dir / f"{name_slug}_enrich.json",
            ])
    for path in candidates:
        if path.exists():
            return json.loads(path.read_text()), str(path)
    return {}, ""


def enrich_value(data: Any, *keys: str) -> str:
    current = data
    for key in keys:
        if not isinstance(current, dict):
            return ""
        current = current.get(key)
    return clean(current)


def build_observations(row: dict[str, str], enrich: dict[str, Any], enrich_path: str) -> str:
    parts = [
        f"Fonte: {clean(row.get('fonte'))}" if clean(row.get("fonte")) else "",
        f"Evidência: {clean(row.get('evidência'))}" if clean(row.get("evidência")) else "",
        f"Hipótese de dor: validar sinais compatíveis com {clean(row.get('ICP'))}." if clean(row.get("ICP")) else "",
        f"Oferta sugerida: {clean(row.get('oferta sugerida'))}" if clean(row.get("oferta sugerida")) else "",
        f"Próxima ação: {clean(row.get('próximo passo'))}" if clean(row.get("próximo passo")) else "",
        f"Classe/score: {clean(row.get('classe'))} | {clean(row.get('score'))}" if clean(row.get("classe")) or clean(row.get("score")) else "",
        f"Contato base: {clean(row.get('contato'))}" if clean(row.get("contato")) else "",
        f"Link e-mail: {clean(row.get('link_email_pronto'))}" if clean(row.get("link_email_pronto")) else "",
        f"Link WhatsApp: {clean(row.get('link_whatsapp_pronto'))}" if clean(row.get("link_whatsapp_pronto")) else "",
        f"Artefato enrich: {enrich_path}" if enrich_path else "",
        f"Pendências: {clean(row.get('risco/pendência'))}" if clean(row.get("risco/pendência")) else "",
    ]
    mailto = enrich_value(enrich, "link_email_pronto")
    whatsapp = enrich_value(enrich, "link_whatsapp_pronto")
    if mailto and not clean(row.get("link_email_pronto")):
        parts.append(f"Link e-mail: {mailto}")
    if whatsapp and not clean(row.get("link_whatsapp_pronto")):
        parts.append(f"Link WhatsApp: {whatsapp}")
    return "\n".join(part for part in parts if part)


def build_payload(row: dict[str, str], enrich: dict[str, Any], enrich_path: str, min_score: float) -> dict[str, str]:
    city, state = parse_city_state(row.get("cidade/UF", ""))
    person_name, contact_email, phone = contact_parts(row.get("contato", ""))
    company_name = clean(row.get("lead"))
    website = normalize_website(clean(row.get("site")) or enrich_value(enrich, "company", "website") or enrich_value(enrich, "website"))
    responsible_email = contact_email or clean(row.get("E-mail do Responsável"))
    markers = join_unique([
        "MEJ/Lead" if "ej" in normalize(row.get("ICP")) or "empresa júnior" in normalize(row.get("lead")) else "Lead",
        clean(row.get("federação")),
        clean(row.get("classe")),
        clean(row.get("ICP")),
    ])
    relation = join_unique([
        "Prospecção OrganizeJr",
        clean(row.get("camada")),
    ], sep=" | ")
    return {
        "Nome - Pessoa": person_name,
        "Cargo - Pessoa": clean(row.get("cargo")),
        "Departamento - Pessoa": clean(row.get("departamento")),
        "CPF - Pessoa": digits(row.get("cpf")),
        "Data de nascimento - Pessoa": clean(row.get("data de nascimento")),
        "Telefones - Pessoa": phone,
        "E-mail - Pessoa": responsible_email if person_name else "",
        "Razão Social - Empresa": company_name,
        "Segmento de atuação - Empresa": derive_segment(row),
        "Origem - Empresa": clean(row.get("origem")) or "Prospecção",
        "CNPJ - Empresa": digits(row.get("CNPJ")),
        "Marcadores - Empresa": markers,
        "Estado - Empresa": state,
        "Cidade - Empresa": city,
        "Site": website,
        "E-mail do Responsável": responsible_email,
        "Relação": relation,
        "Observações": build_observations(row, enrich, enrich_path),
        "Enviar ao Ploomes?": classify_send_flag(row, min_score),
    }


def read_cell(ws, row_idx: int, header_map: dict[str, int], field: str) -> str:
    col = header_map.get(field)
    if not col:
        return ""
    return clean(ws.cell(row=row_idx, column=col).value)


def match_row(ws, header_map: dict[str, int], payload: dict[str, str], start_row: int) -> int | None:
    new_cnpj = digits(payload.get("CNPJ - Empresa"))
    new_name = normalize(payload.get("Razão Social - Empresa"))
    new_city = normalize(payload.get("Cidade - Empresa"))
    new_state = normalize(payload.get("Estado - Empresa"))
    new_email = normalize(payload.get("E-mail do Responsável"))
    for row_idx in range(start_row + 1, ws.max_row + 1):
        existing_cnpj = digits(read_cell(ws, row_idx, header_map, "CNPJ - Empresa"))
        if new_cnpj and existing_cnpj == new_cnpj:
            return row_idx
        existing_name = normalize(read_cell(ws, row_idx, header_map, "Razão Social - Empresa"))
        existing_city = normalize(read_cell(ws, row_idx, header_map, "Cidade - Empresa"))
        existing_state = normalize(read_cell(ws, row_idx, header_map, "Estado - Empresa"))
        existing_email = normalize(read_cell(ws, row_idx, header_map, "E-mail do Responsável"))
        if new_email and existing_email == new_email and new_name and existing_name == new_name:
            return row_idx
        if new_name and existing_name == new_name and new_city and existing_city == new_city and new_state and existing_state == new_state:
            return row_idx
    return None


def set_field(ws, row_idx: int, header_map: dict[str, int], field: str, value: str) -> None:
    col = header_map.get(field)
    if not col:
        return
    ws.cell(row=row_idx, column=col).value = value


def write_payload(ws, row_idx: int, header_map: dict[str, int], payload: dict[str, str], is_existing: bool) -> None:
    for field, value in payload.items():
        if field in CONTROL_HEADERS and field != "Enviar ao Ploomes?":
            continue
        if field == "Enviar ao Ploomes?" and is_existing:
            current = normalize(read_cell(ws, row_idx, header_map, field))
            if current in {"não", "nao", "no"}:
                continue
            if normalize(read_cell(ws, row_idx, header_map, "Status Importação")) == "enviado":
                continue
        set_field(ws, row_idx, header_map, field, value)


def process(args: argparse.Namespace) -> dict[str, Any]:
    lead_table, enrichment_dir, workbook_path = resolve_paths(args)
    rows = read_table(lead_table, args.limit)
    wb = load_workbook(workbook_path)
    ws = wb[wb.sheetnames[0]]
    header_row, headers = find_header_row(ws)
    header_map = build_header_map(headers)
    summary = {
        "lead_table": str(lead_table),
        "workbook": str(workbook_path),
        "output": str(workbook_path if args.in_place else Path(args.output).expanduser().resolve()),
        "rows_read": len(rows),
        "matched": 0,
        "appended": 0,
        "updated": 0,
        "send_sim": 0,
        "send_nao": 0,
        "sample": [],
    }
    for row in rows:
        enrich, enrich_path = load_enrichment(row, enrichment_dir)
        payload = build_payload(row, enrich, enrich_path, args.min_score)
        row_idx = match_row(ws, header_map, payload, header_row)
        is_existing = row_idx is not None
        if row_idx is None:
            row_idx = ws.max_row + 1
            summary["appended"] += 1
        else:
            summary["matched"] += 1
            summary["updated"] += 1
        write_payload(ws, row_idx, header_map, payload, is_existing)
        if payload["Enviar ao Ploomes?"] == "SIM":
            summary["send_sim"] += 1
        else:
            summary["send_nao"] += 1
        if len(summary["sample"]) < 5:
            summary["sample"].append({
                "lead": payload["Razão Social - Empresa"],
                "row": row_idx,
                "send": payload["Enviar ao Ploomes?"],
                "matched_existing": is_existing,
            })
    out_path = workbook_path if args.in_place else Path(args.output).expanduser().resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return summary


def main() -> int:
    args = parse_args()
    summary = process(args)
    if args.json:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
    else:
        print(f"rows_read={summary['rows_read']}")
        print(f"matched={summary['matched']}")
        print(f"appended={summary['appended']}")
        print(f"updated={summary['updated']}")
        print(f"send_sim={summary['send_sim']}")
        print(f"send_nao={summary['send_nao']}")
        print(f"output={summary['output']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
